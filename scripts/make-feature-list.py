"""Generate FatigueIDPro_Feature_List.docx - a complete, categorised list of every
software feature, split into Participant and Admin sections. Run: python scripts/make-feature-list.py
"""
from datetime import date
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x0F, 0x2B, 0x4A)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GREY = RGBColor(0x47, 0x55, 0x69)

doc = Document()

# base styles
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)

def title(text, sub):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = NAVY
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(sub); r2.font.size = Pt(12); r2.font.color.rgb = BLUE
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Feature reference  |  " + date.today().strftime("%d %B %Y"))
    r3.font.size = Pt(9); r3.font.color.rgb = GREY

def section(text):
    doc.add_paragraph()
    h = doc.add_heading(level=1)
    r = h.add_run(text); r.font.color.rgb = NAVY; r.font.size = Pt(16)

def subsection(text):
    h = doc.add_heading(level=2)
    r = h.add_run(text); r.font.color.rgb = BLUE; r.font.size = Pt(12.5)

def feature(name, desc):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(name + " — "); r.bold = True; r.font.color.rgb = NAVY
    p.add_run(desc)

def para(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.color.rgb = GREY; r.font.size = Pt(10)

# ---------------------------------------------------------------- cover
title("FatigueIDPro", "Complete Software Feature List — Participant Experience & Research Admin Console")
para("A web-based platform for collecting human-fatigue research data. Participants take a "
     "standardised, randomised test battery in the browser; researchers manage accounts, review "
     "data quality, and export the collected data from a secure admin console. Data is stored in a "
     "Supabase (PostgreSQL) backend with row-level security. This document lists every feature.")

# ================================================================ PARTICIPANT
section("Part A — Participant Features")

subsection("Entry & Informed Consent")
feature("Single shared entry", "One URL with a Participant / Researcher chooser — no separate links to manage.")
feature("Informed consent screen", "IRB-style consent showing institution, protocol ID, data use, retention period, contact, participant rights (voluntary, withdraw anytime), risks, and data handling. The participant must explicitly agree before anything is collected.")
feature("Configurable study details", "Institution, contact, protocol ID, retention, data-use and physical-protocol text are set per study; unfinalised fields show 'To be confirmed'.")
feature("Consent decline path", "Declining ends cleanly with no data collected.")

subsection("Camera Attention & Calibration")
feature("Optional camera opt-in", "A separate, clearly worded camera-consent screen; the participant can take part fully without the camera.")
feature("Explicit photo-capture opt-in", "A dedicated tick-box consents to saving occasional low-resolution photos; unticking keeps camera use on-device only.")
feature("On-device attention tracking", "MediaPipe FaceLandmarker estimates 'looking at screen / away' entirely on the device (no video leaves the browser for the metric).")
feature("Adaptive 4-dot calibration", "Four dots appear at the true screen corners; the head angle at each corner defines that person's comfortable viewing range, which becomes their personal attention threshold — adapting to distance and screen size.")
feature("Calibration ground-truth photos", "If photo capture is consented, one labelled photo is saved at each corner (top-left/right, bottom-left/right) as reference for how the person looks at each part of the screen.")
feature("App-switch & away detection", "Detects when the participant leaves the tab/app and how long they were away, per task.")

subsection("Registration & Demographics")
feature("Self-registration with auto-ID", "Participant enters an email (optionally name/phone); the app issues a unique participant ID automatically.")
feature("Duplicate protection", "Registration is de-duplicated by email — the same person always gets the same ID, so one person cannot enrol twice.")
feature("ID confirmation screen", "The assigned ID is shown so the participant can save it (for questions or withdrawal).")
feature("Separated contact storage", "Email/name/phone are stored apart from the research data and are visible only to an admin; research tables stay pseudonymous.")
feature("Demographics form", "Age, gender, primary input device, dominant hand, and eye correction.")

subsection("The Test Battery (3 blocks)")
feature("Reproducible randomised protocol", "Base task (Fitts or Typing) and fatigue track (Cognitive or Physical) are assigned from a per-participant seed, so the exact sequence is reproducible.")
feature("KSS sleepiness rating", "Karolinska Sleepiness Scale captured before and after each block (a validated instrument).")
feature("Fitts' Law pointing task", "Target acquisition measuring movement time, error rate and throughput.")
feature("Typing / transcription task", "Measures typing speed, accuracy and inter-key timing.")
feature("NASA-TLX workload", "Six-subscale subjective workload rating (mental, physical, temporal, performance, effort, frustration).")
feature("Cognitive task (Stroop / AX-CPT)", "Reaction-time and accuracy on congruent/incongruent and cue/probe trials.")
feature("Physical fatigue task", "A researcher-approved movement protocol with a safety screening and an alternative for anyone with a health concern.")

subsection("Progress & Experience")
feature("Overall completion percentage", "A live 'Test completion' bar and percentage (0-100%) so the participant always knows how much of the test remains.")
feature("Protocol stepper", "A four-stage stepper (Consent - Setup - Data Collection - Results) plus the current Block X/3.")
feature("Fullscreen kiosk mode", "The session runs full screen like an online-assessment; leaving full screen shows a blocking overlay that must be resumed before continuing.")
feature("Smooth, modern UI", "Animated, glassmorphic interface with clear task instructions.")
feature("Accessibility", "Respects reduced-motion, uses ARIA progress roles and labelled form fields.")

subsection("Reliability & Safety")
feature("Offline write-queue", "Records are queued and retried if the network drops, with idempotency so nothing is saved twice — no data loss on a bad connection.")
feature("Dead-letter tracking", "Permanently failed writes are recorded rather than silently lost.")
feature("Resume interrupted session", "If the browser closes mid-test, the participant can resume from where they stopped.")
feature("Accidental-exit protection", "Warns before an accidental refresh or tab-close during the test.")
feature("Withdraw at any time", "A always-available Withdraw button stops all further collection and marks the session abandoned (voluntary participation).")
feature("Truthful completion status", "Sessions are explicitly finalised (status + completion time) so analysts can filter genuinely finished runs.")
feature("Device/context capture", "Screen size, input method, pixel ratio and orientation are recorded for analysis.")

# ================================================================ ADMIN
section("Part B — Admin / Research Console Features")

subsection("Access & Security")
feature("Secure researcher login", "Email/password authentication (Supabase Auth).")
feature("Approval-gated sign-up", "New accounts are created as 'pending' with no data access; they see an 'awaiting approval' screen until an admin approves them.")
feature("Access-request queue (admin)", "Admins approve or deny each account and assign a role (viewer / researcher / admin).")
feature("Role-based access", "Three roles enforced by the database itself (row-level security), not just the UI.")
feature("Append-only audit log", "Every approval, denial and data export is recorded (who, what, when, row counts) in a tamper-resistant, admin-only log.")

subsection("Dashboard & Sessions")
feature("Multi-section console", "A sidebar-navigated console (not one long page), with a calm light theme and subtle animated background.")
feature("Overview dashboard", "KPI cards (total / completed / in-progress sessions, IDs available / assigned, measurements) with count-up animation, recent sessions and quick links.")
feature("Sessions browser", "A filterable list of all sessions; click one to open its full detail.")
feature("Per-session detail", "Status, protocol, average NASA-TLX and camera attention, app-switches and time away; a per-test engagement table; uploaded measurements; and a camera-snapshot gallery (via short-lived signed links).")

subsection("Data Quality & Review")
feature("Data-quality go/no-go", "Each session is flagged Good / Review / Incomplete based on completion and app-visibility, with attention %, app-switches and time-away.")
feature("Exploratory-attention guardrail", "Camera attention is clearly labelled exploratory and must not be used as an automatic exclusion criterion.")

subsection("Per-Test Data Browsers")
feature("Dedicated page per test", "Separate, filterable data pages for Fitts, Typing, Cognitive, NASA-TLX, Fatigue (KSS), Attention/Engagement, Scroll sessions and Scroll intervals.")
feature("Live filter & row counts", "Client-side search hides non-matching rows with a live count on every table.")
feature("Per-dataset CSV", "One-click CSV download for the specific dataset being viewed (audited).")

subsection("Participants, Measurements & Export")
feature("Registrations view (admin)", "Lists self-registered participants with their contact details; filter and export (admin-only).")
feature("Participant ID pool", "Generate batches of participant codes, revoke a withdrawn code, or release a code for reuse (for researcher-issued studies).")
feature("Manual measurement upload", "Upload an ECG (heart rate, HRV) or a manual physical result for a participant/session; view the list and export it.")
feature("CSV export centre", "Export any single dataset or all datasets at once (RFC-4180 CSV for Excel / pandas / R); every export is written to the audit log.")
feature("Live sidebar counts", "Session and available-ID counts shown in the navigation.")
feature("Responsive layout", "The console collapses to a mobile-friendly layout on small screens.")

# ================================================================ BACKEND
section("Part C — Security & Data Handling (Backend)")
feature("PostgreSQL with row-level security", "Every table is protected by RLS; participants can only insert their own data, and only approved researchers can read it.")
feature("Pseudonymous research data", "Research tables reference a participant code only; personal contact details live in a separate, admin-only table.")
feature("Private image storage", "Camera snapshots go to a private storage bucket, readable only by an admin through expiring signed links.")
feature("Reproducible & versioned", "Per-participant randomisation seeds and versioned metric definitions are stored with each session.")
feature("Global camera kill-switch", "Image collection can be turned off entirely with one environment setting (on-device metrics still work) to control storage cost.")
feature("Research export views", "Clean, pseudonymous export views power the CSV downloads.")
feature("Version-controlled migrations", "All database changes are captured as migrations for reproducibility.")

para("Note: A separate optional 'scroll study' module (Instagram-style scroll-fatigue measurement) also exists in the codebase and shares the same backend.")

out = "FatigueIDPro_Feature_List.docx"
doc.save(out)
print("wrote", out)
