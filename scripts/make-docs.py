#!/usr/bin/env python3
"""
Generate the two hand-over Word documents for FatigueIDPro:

  FatigueIDPro_User_Guide.docx   - how to set up, operate, and run sessions
  FatigueIDPro_Features.docx     - every feature and how it helps the research

Re-run after changes to keep them current:  python scripts/make-docs.py
Requires: pip install python-docx
"""
from docx import Document
from docx.shared import Pt, RGBColor

ROOT = "e:/Data-center/fatigue-analysis"
BLUE = RGBColor(0x1E, 0x40, 0xAF)


def build(path, blocks):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for kind, text in blocks:
        if kind == "title":
            doc.add_heading(text, level=0)
        elif kind == "h1":
            h = doc.add_heading(text, level=1)
            for r in h.runs:
                r.font.color.rgb = BLUE
        elif kind == "h2":
            doc.add_heading(text, level=2)
        elif kind == "b":
            doc.add_paragraph(text, style="List Bullet")
        elif kind == "n":
            doc.add_paragraph(text, style="List Number")
        else:
            doc.add_paragraph(text)
    doc.save(path)
    print("wrote", path)


# ---------------------------------------------------------------- USER GUIDE
user_guide = [
    ("title", "FatigueIDPro - User Guide"),
    ("p", "FatigueIDPro is a research platform that measures fatigue from how a person interacts with a "
          "device. This guide explains how to set it up, run a data-collection session, review the results, "
          "and manage researcher access."),

    ("h1", "1. What you need"),
    ("b", "A free Supabase project (the database + secure file storage)."),
    ("b", "The app, hosted at a URL you deploy (e.g. Vercel) or run locally with npm."),
    ("b", "An approved researcher (admin) account to use the console."),

    ("h1", "2. One-time setup"),
    ("n", "In Supabase, open the SQL Editor and run backend/supabase/schema.sql once, then apply the files "
          "in backend/supabase/migrations in date order. This creates all tables, security rules, views, "
          "functions, and the private image-storage bucket."),
    ("n", "Set the environment variables (in Vercel, or a local .env): VITE_SUPABASE_URL and "
          "VITE_SUPABASE_ANON_KEY. Also set the two required study details shown on the consent screen: "
          "VITE_STUDY_INSTITUTION and VITE_STUDY_CONTACT. The other VITE_STUDY_* values (protocol ID, "
          "retention, physical-protocol text) are optional and show 'To be confirmed' until filled."),
    ("n", "Create the FIRST admin: in Supabase, Authentication > Users > Add user, then in the SQL Editor "
          "insert a row in researcher_profiles for that user with role 'admin'. (After this, further "
          "researchers sign themselves up and you approve them in the console - see section 4.)"),
    ("n", "Deploy. On Vercel, pushing to the main branch builds automatically; otherwise Redeploy after any "
          "change or environment-variable update."),

    ("h1", "3. The link and the two roles"),
    ("p", "There is one link to share. The home page asks whether the person is a Participant or a Researcher:"),
    ("b", "Home (/): 'I'm a Participant' starts the study; 'Researcher / Admin' opens the console."),
    ("b", "/admin.html: the research console (sign-in required)."),
    ("b", "/scroll.html: the phone-use study (runs fully only inside the Android app - see section 10)."),

    ("h1", "4. Getting and granting researcher access"),
    ("n", "A researcher opens the console and creates an account (email + password)."),
    ("n", "New accounts are 'pending' and can see NO data until approved. They see an 'awaiting approval' "
          "screen."),
    ("n", "An admin opens Administration > Access requests, and Approves the account with a role "
          "(viewer, researcher, or admin) or Denies it."),
    ("p", "Every approval, denial, and data export is written to an append-only Audit log (Administration > "
          "Audit log) - who did what, and when."),

    ("h1", "5. Running a participant session"),
    ("n", "Open the link on their device and tap 'I'm a Participant'."),
    ("n", "They read and agree to the informed-consent screen. The session then goes full screen (like an "
          "online exam); leaving full screen shows a 'return to full screen' prompt."),
    ("n", "They can optionally enable the camera for an attention check, and separately tick whether "
          "occasional low-resolution photos may be saved. If the camera is on, a 4-dot calibration appears - "
          "they look at and tap each corner dot (this adapts the attention check to their screen and "
          "distance)."),
    ("n", "They register: they enter their email (and optionally name/phone). The app issues a participant "
          "ID automatically and shows it to save. The same email always gets the same ID, so no one can "
          "enrol twice."),
    ("n", "They complete three blocks. Each block: a sleepiness (KSS) rating, a tapping or typing task, a "
          "NASA-TLX workload rating, and a cognitive or light physical task, then a KSS rating again."),
    ("n", "A 'Test completion' bar shows their overall progress (0-100%). At the end they see a summary."),
    ("p", "Notes: a Withdraw button is always available. If the page is refreshed or the connection drops, "
          "the session resumes where it left off and any data that could not upload is retried automatically "
          "- nothing is lost."),

    ("h1", "6. Uploading manual measurements (ECG / physical)"),
    ("n", "In the console, open Operations > Measurements."),
    ("n", "Choose the participant's session, pick ECG or Physical, enter the values, and upload."),
    ("p", "They appear in that participant's session detail. (Automatic ECG upload from a Raspberry Pi is a "
          "planned addition.)"),

    ("h1", "7. Reviewing results and data quality"),
    ("b", "Overview: totals at a glance (sessions, completed, IDs, measurements) and recent sessions."),
    ("b", "Sessions: click a session to see its per-test results, attention validation, uploaded "
          "measurements, and - if the camera was on - a gallery of the saved snapshots."),
    ("b", "Data quality: a go/no-go flag per session - Good, Review (left the app often), or Incomplete."),
    ("b", "Test data: a page per test (Fitts, Typing, Cognitive, NASA-TLX, Fatigue/KSS, Attention, and the "
          "phone-usage data), each filterable with its own CSV download."),
    ("b", "Administration > Registrations: self-registered participants and their contact details "
          "(admin-only)."),
    ("b", "Export: download any dataset, or all of them, as CSV for Excel, Python (pandas), or R."),

    ("h1", "8. Camera images and storage"),
    ("p", "If a participant consents to photo capture, the app saves low-resolution photos to a private "
          "storage bucket (four labelled calibration photos, plus about one every 15 seconds during the "
          "test). Only an admin can view them. To turn image collection off entirely (on-device attention "
          "metrics still work), set VITE_COLLECT_CAMERA_FRAMES=false and redeploy."),

    ("h1", "9. Participant ID pool (optional)"),
    ("p", "Self-registration is the default. If you instead want to hand out researcher-issued codes, use "
          "Operations > Participant IDs to generate a batch, and Revoke or Release codes as needed."),

    ("h1", "10. The phone-use study"),
    ("p", "This study measures real phone use (Instagram, Facebook, YouTube, etc.): how long each app is "
          "used, how often it is opened, and how much the participant scrolls, and how that drifts over 30, "
          "60, or 120 minutes. Because phones sandbox apps, this needs the FatigueIDPro Android app (a "
          "browser cannot see other apps). The web page handles registration and the KSS ratings; the app "
          "does the background recording. See the separate 'Phone-Usage Study App' note for architecture, "
          "permissions, and why it is distributed as a direct download rather than via the Play Store."),

    ("h1", "11. Troubleshooting"),
    ("b", "'Awaiting approval' after signing in: an admin must approve the account in Access requests."),
    ("b", "New features not showing on the live site: redeploy, then hard-refresh (Ctrl+Shift+R)."),
    ("b", "Consent screen shows 'Study setup is incomplete': set VITE_STUDY_INSTITUTION and "
          "VITE_STUDY_CONTACT and redeploy."),
    ("b", "'Nothing is saving': check the Supabase URL/key and that schema.sql + migrations have been run."),
]

# ------------------------------------------------------------- FEATURES DOC
features = [
    ("title", "FatigueIDPro - Features and How They Help"),
    ("p", "This document explains each part of FatigueIDPro and why it matters for collecting reliable "
          "fatigue-research data."),

    ("h1", "Interaction tasks and KSS ratings"),
    ("p", "Fitts tapping, typing, NASA-TLX workload, a cognitive battery (Stroop and AX-CPT), and a light "
          "physical task run across three blocks, with a Karolinska Sleepiness Scale (KSS) rating before and "
          "after each block."),
    ("b", "How it helps: fatigue shows up as performance change over time; combining fine-grained behaviour "
          "with a validated subjective scale gives a rich, trusted signal for analysis and machine learning."),

    ("h1", "Reproducible, randomized assignment"),
    ("p", "Each session's task order is randomly assigned from a stored per-participant seed, and every "
          "metric carries a version tag."),
    ("b", "How it helps: prevents selection bias, keeps the design balanced, and makes the exact sequence "
          "reproducible for auditing and re-analysis."),

    ("h1", "Self-registration with automatic, unique IDs"),
    ("p", "Participants register with their email; the app issues a unique participant ID automatically and "
          "de-duplicates by email. Contact details are stored in a separate, admin-only table, apart from the "
          "research data."),
    ("b", "How it helps: anyone can take part without a researcher handing out codes, no one can enrol "
          "twice, and the research data stays pseudonymous."),

    ("h1", "Consent, camera opt-in, safety screening, and withdrawal"),
    ("p", "Informed consent up front; a separate camera opt-in with an explicit tick-box for saving photos; "
          "a physical-activity safety check; and an always-available Withdraw button."),
    ("b", "How it helps: responsible, ethics-aligned research with a clear record of what each person agreed "
          "to."),

    ("h1", "Full-screen kiosk mode"),
    ("p", "The session runs full screen like an online exam; leaving full screen shows a blocking prompt to "
          "return before continuing."),
    ("b", "How it helps: keeps attention on the tasks and reduces distraction, improving data quality."),

    ("h1", "Adaptive attention validation with optional snapshots"),
    ("p", "An on-device camera check estimates whether the participant is looking at the screen. A 4-dot "
          "calibration measures their comfortable viewing angles and adapts the threshold to their screen and "
          "distance. If they consent, low-resolution photos are saved to private storage - four labelled "
          "calibration 'ground-truth' photos plus periodic photos during the test - and image collection has "
          "a single global on/off switch."),
    ("b", "How it helps: you can trust each measurement and flag distracted sessions; the calibration photos "
          "give a per-screen reference for studying attention across different devices, while a kill-switch "
          "controls storage cost and privacy."),

    ("h1", "Live completion progress"),
    ("p", "A 'Test completion' bar shows the participant how far through the test they are (0-100%)."),
    ("b", "How it helps: sets expectations and reduces drop-off part way through."),

    ("h1", "Session resume and offline write-queue"),
    ("p", "A refresh or dropped connection restores the exact step; any data that fails to upload is parked "
          "and retried, and permanent failures are recorded rather than lost."),
    ("b", "How it helps: no participant restarts, and no data point is lost on unreliable venue Wi-Fi."),

    ("h1", "Session finalization"),
    ("p", "Sessions are marked completed or abandoned with a timestamp."),
    ("b", "How it helps: analysts can filter for sessions that genuinely finished instead of guessing."),

    ("h1", "Multi-section research console"),
    ("p", "A sidebar console (light, mobile-friendly) with an overview dashboard, a sessions browser with "
          "per-session detail and a snapshot gallery, a data-quality view, a page per test with its own CSV, "
          "a registrations list, participant-ID tools, and manual measurement upload."),
    ("b", "How it helps: fast monitoring during collection and a quick way to decide which data to keep."),

    ("h1", "Researcher access control and audit log"),
    ("p", "Sign-ups start with no access; an admin approves each account and assigns a role (viewer, "
          "researcher, admin), enforced by the database itself. Every approval, denial, and data export is "
          "written to an append-only audit log."),
    ("b", "How it helps: only authorised people ever see data, and there is a tamper-resistant record of who "
          "did what - important for a research institution."),

    ("h1", "Data-quality go/no-go review"),
    ("p", "Each session is flagged Good, Review, or Incomplete from completion and app-visibility signals; "
          "camera attention is treated as exploratory, not an automatic exclusion."),
    ("b", "How it helps: a quick, consistent basis for deciding which sessions to include."),

    ("h1", "Manual and device measurements (ECG / physical)"),
    ("p", "Researchers upload ECG or manual physical results per participant. Automatic ECG upload from a "
          "Raspberry Pi is planned."),
    ("b", "How it helps: combines objective physiological data with the behavioural measures for a fuller "
          "picture of fatigue."),

    ("h1", "Phone-use (real app usage) study"),
    ("p", "Measures fatigue from real phone use: per app (Instagram, Facebook, YouTube, ...) how long it is "
          "used, how often it is opened, and how much the participant scrolls, and how that drifts over the "
          "session. It runs in a native Android app that records in the background; no content is captured."),
    ("b", "How it helps: captures fatigue from natural, everyday phone use rather than an artificial in-app "
          "feed - a far more realistic signal."),

    ("h1", "One-click CSV export"),
    ("p", "Download any dataset, or all of them, as CSV from the console (each export is logged)."),
    ("b", "How it helps: the data goes straight into Excel, pandas, or R - no manual database queries."),

    ("h1", "Secure, centralized data"),
    ("p", "All data lives in one Supabase database with row-level security: the public app can only insert, "
          "only approved researchers can read, contact details and camera images are locked to admins, and "
          "all schema changes are version-controlled migrations."),
    ("b", "How it helps: no fragile manual file collection, and participant data is protected by "
          "least-privilege access."),
]

build(ROOT + "/FatigueIDPro_User_Guide.docx", user_guide)
build(ROOT + "/FatigueIDPro_Features.docx", features)
