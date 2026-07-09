"""Generate FatigueIDPro_DL_Roadmap.docx - the fatigue DL model plan: goal/phases,
per-dataset labels, model architecture, and how it helps collection + model building.
Run: python scripts/make-dl-doc.py
"""
from datetime import date
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x0F, 0x2B, 0x4A)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GREY = RGBColor(0x47, 0x55, 0x69)

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)


def title(t, sub):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(sub); r2.font.size = Pt(12); r2.font.color.rgb = BLUE
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Planning note (no training yet)  |  " + date.today().strftime("%d %B %Y"))
    r3.font.size = Pt(9); r3.font.color.rgb = GREY


def h1(t):
    doc.add_paragraph()
    h = doc.add_heading(level=1); r = h.add_run(t); r.font.color.rgb = NAVY; r.font.size = Pt(15)


def h2(t):
    h = doc.add_heading(level=2); r = h.add_run(t); r.font.color.rgb = BLUE; r.font.size = Pt(12)


def para(t, grey=False):
    p = doc.add_paragraph(); r = p.add_run(t)
    if grey: r.font.color.rgb = GREY; r.font.size = Pt(10)


def bullet(name, desc):
    p = doc.add_paragraph(style="List Bullet")
    if name:
        r = p.add_run(name + " - "); r.bold = True; r.font.color.rgb = NAVY
    p.add_run(desc)


def table(rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.text = val
            for pp in c.paragraphs:
                for rn in pp.runs:
                    rn.font.size = Pt(8.5)
                    if ri == 0: rn.bold = True


# -------------------------------------------------------------- cover
title("FatigueIDPro - Fatigue Model Roadmap",
      "Per-dataset labels, lightweight/edge model, and how it supports data collection")
para("Goal: a lightweight, edge-deployable fatigue model - benchmarked on public datasets, "
     "then refined on our own multimodal data. Small enough to run on a microcontroller. This "
     "note covers the plan, the label each dataset gives us, and how the model helps while we "
     "collect data and as we keep improving it. No training is done yet.", grey=True)

# -------------------------------------------------------------- phases
h1("1. Phases")
bullet("P0 - Benchmark on public data", "Fixed features (HR, HRV, reaction-time/error drift) + "
       "classic baselines (logistic regression, gradient-boosted trees) to set a reference bar.")
bullet("P1 - Lightweight model", "Small 1D-CNN / tiny temporal model on windowed signals; quantise "
       "to int8 and export to TensorFlow Lite Micro / ONNX for microcontroller deployment.")
bullet("P2 - Train on OUR data", "Same pipeline on our labelled sessions (tasks + NASA-TLX + Borg + "
       "KSS + ECG/EMG); compare against P0/P1 to show our data adds value.")
bullet("P3 - Advanced deep learning", "Multimodal fusion + sequence models. Deferred until requirements "
       "are clarified.")

# -------------------------------------------------------------- labels
h1("2. Labels per dataset")
para("We do NOT force one label across datasets. Each dataset type keeps the label that suits its "
     "signal; we train a model per modality from those. Our own collected data is a NEW multimodal "
     "type that carries every label and ties them together.")
table([
    ["Dataset", "Signals", "Label we use", "What it teaches the model"],
    ["DROZY", "ECG, EEG, EOG, EMG, NIR video", "KSS sleepiness (1-9), PVT lapses", "physiology + face -> sleepiness"],
    ["FatigueSet", "ECG, PPG, EEG, GSR, ST, ACC", "mental-fatigue self-report", "wearable signals -> mental fatigue"],
    ["Fatigue-Characterization (MR)", "EEG, ECG, EDA, SpO2, resp, temp", "fatigue state", "multimodal physiology -> fatigue"],
    ["MEFAR", "EEG, HR, PPG, GSR, ST, ACC", "occupational mental fatigue", "office / mental fatigue"],
    ["UL-DD", "RGB/IR/3D video + SpO2, BVP, HR", "drowsy vs alert", "face + physiology -> drowsiness"],
    ["UTA-RLDD", "RGB webcam video", "alert / low-vigilant / drowsy", "webcam face -> drowsiness (matches our camera)"],
    ["NTHU-DDD", "IR video", "drowsy/not + blink, yawn, eye, head", "facial micro-signs of drowsiness"],
    ["Mendeley EMG (biceps/triceps)", "sEMG", "muscle fatigue (MDF/MNF decline)", "EMG spectral shift -> muscle fatigue"],
    ["Handgrip force-time", "grip force", "fatigue index / endurance decline", "force decline -> physical fatigue"],
    ["MIMIC-III", "ECG, ABP (BP), PPG, SpO2, resp", "none (not fatigue-labelled)", "pretraining + signal-quality; BP/SpO2 baselines"],
    ["Our data (NEW)", "Typing/Pointing, Stroop, grip + ECG + EMG", "KSS, Borg CR10, NASA-TLX, performance drift", "the target multimodal set with all labels"],
])

# -------------------------------------------------------------- architecture
h1("3. How the datasets combine")
bullet("Per-modality models", "an ECG/HRV model (DROZY, FatigueSet, MIMIC), an EMG model (Mendeley "
       "EMG, grip), a face/video model (UTA-RLDD, NTHU, UL-DD, DROZY).")
bullet("Our data = fusion + validation set", "our sessions have every modality plus KSS/Borg/NASA, so "
       "they fuse the per-modality models and calibrate them to our own subjective + behavioural labels.")
bullet("Data-quality label", "a derived valid/artefact label runs on every signal for live verification "
       "during collection.")

# -------------------------------------------------------------- helps: collection
h1("4. How the model helps DURING data collection")
bullet("Live quality gate", "watches incoming ECG + behaviour in real time; flags artefacts, a loose/off "
       "electrode, or a disengaged participant immediately - so we fix it, not discover ruined data later.")
bullet("Self-report cross-check", "compares its fatigue estimate against the participant's KSS/Borg/NASA "
       "answers; a big mismatch flags a careless responder or a mislabelled sample.")
bullet("Manipulation check", "confirms whether the fatigue task actually raised fatigue across the session "
       "- you know on the spot if the design worked.")
bullet("Live dashboard", "operator sees a running fatigue estimate + confidence and catches sensor drops "
       "or drift as they happen.")

# -------------------------------------------------------------- helps: building
h1("5. How the model helps BUILD the model (the loop)")
bullet("Bootstrap from public data", "train the benchmark on the public datasets first, so there is a "
       "working model before we have much of our own data; then fine-tune on ours (transfer learning).")
bullet("Smart labelling (active learning)", "the model pre-labels each new session; the operator only "
       "reviews the uncertain ones - fast, cheap labelling of our growing dataset.")
bullet("Signal importance", "comparing per-modality models shows which signals actually carry fatigue "
       "information, so we invest sensor effort where it pays off.")
bullet("Continuous improvement + proof", "every validated session -> retrain -> accuracy climbs; benchmark "
       "against the public-only model to prove our data adds value.")
bullet("Edge deployment", "quantise the final model to a microcontroller: a standalone wearable fatigue "
       "monitor, and on-device verification with no cloud.")
para("Short version: the same model is both a quality inspector while we collect and the research "
     "deliverable we keep improving. Each session makes it better; a better model makes the next "
     "collection cleaner.", grey=True)

doc.save("FatigueIDPro_DL_Roadmap.docx")
print("wrote FatigueIDPro_DL_Roadmap.docx")
