"""Generate dl-model/send/FatigueIDPro_Data_Collection_Abstract.docx - a plain-
language summary of every category of data collected from a participant,
verified against the live database schema. Run:
    python scripts/make-data-collection-abstract-doc.py
"""
from datetime import date
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x0F, 0x2B, 0x4A); BLUE = RGBColor(0x1D, 0x4E, 0xD8); GREY = RGBColor(0x47, 0x55, 0x69)
OUT = "e:/Data-center/fatigue-analysis/dl-model/send/FatigueIDPro_Data_Collection_Abstract.docx"

doc = Document()
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10.5)


def title(t, sub):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = True; r.font.size = Pt(19); r.font.color.rgb = NAVY
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(sub); r2.font.size = Pt(11.5); r2.font.color.rgb = BLUE
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Center of Excellence in Data Science and AI, TIET Patiala  |  " + date.today().strftime("%d %B %Y"))
    r3.font.size = Pt(9); r3.font.color.rgb = GREY


def h1(t):
    doc.add_paragraph(); h = doc.add_heading(level=1)
    r = h.add_run(t); r.font.color.rgb = NAVY; r.font.size = Pt(13.5)


def para(t, grey=False, bold=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold
    if grey: r.font.color.rgb = GREY; r.font.size = Pt(10)


def bullet(name, desc=None):
    p = doc.add_paragraph(style="List Bullet")
    if desc is None:
        p.add_run(name)
    else:
        r = p.add_run(name + " - "); r.bold = True; r.font.color.rgb = NAVY
        p.add_run(desc)


title("Participant Data Collection - Abstract",
      "Examining Real-Time Fatigue Through Interaction Analysis")
para("This document summarises every category of data collected from a participant by the "
     "FatigueIDPro platform, in plain terms, verified directly against the live database "
     "schema (not from memory) so it is accurate as of this date.", grey=True)

h1("1. Identity and registration")
bullet("Email (required)", "used only to auto-issue a unique participant ID and to prevent "
       "the same person registering twice. Stored in a separate, admin-only table, never "
       "joined to research data by anyone but an admin.")
bullet("Full name and phone (optional)", "same separate, admin-only table.")
bullet("Auto-issued participant ID", "e.g. FP-A1B2C3 - this pseudonymous code, not the "
       "participant's name or email, is what every research record is tagged with.")

h1("2. Consent and study context")
bullet("A record that informed consent was given, with a timestamp and the protocol version shown.")
bullet("The study details shown at consent (institution, protocol ID, retention period, contact).")
bullet("If the physical-fatigue track applies: answers to the physical activity safety "
       "screening (PAR-Q style) and its outcome.")

h1("3. Demographics")
para("Age, gender, primary input device (mouse / trackpad / touchscreen / other), dominant "
     "hand, and eye correction status (normal / glasses / contacts / uncorrected).")

h1("4. Device and session context")
para("Browser/device type, screen size and pixel ratio, platform, a session identifier, the "
     "randomisation seed that determined task order, and the software build version. "
     "Technical context for analysis, not personal identification.")

h1("5. Camera and attention data (optional, separately consented)")
para("The camera is entirely opt-in and asked about separately from the main consent. If enabled:")
bullet("On-device attention signal", "only a 'looking at the screen' percentage and a "
       "look-away count per task, computed on the participant's own device. No video is "
       "analysed off-device for this signal.")
bullet("4-dot calibration", "screen-tap coordinates and head-angle (yaw/pitch) at each of the "
       "four screen corners, used to adapt the attention threshold to that person's screen "
       "and seating distance.")
bullet("App-switch / tab-away tracking", "how many times and for how long the participant "
       "left the browser tab during a task.")
bullet("Low-resolution photo snapshots (only if separately consented)", "four labelled photos "
       "during the 4-dot calibration, plus roughly one every 15 seconds during the test. "
       "Stored in a private area readable only by an admin. Can be turned off platform-wide "
       "with a single setting, with no effect on the numeric attention signal above.")

h1("6. Self-report questionnaires")
para("Collected at four points per session (once before any task, and after each of the "
     "three main task rounds):")
bullet("NASA-TLX", "six workload ratings (mental demand, physical demand, temporal demand, "
       "performance, effort, frustration), each 1-20.")
bullet("Borg CR10", "a single perceived-exertion rating, 0-10.")
para("Collected at two points per session (the first and last questionnaire only):")
bullet("KSS (Karolinska Sleepiness Scale)", "a single sleepiness rating, 1-9, giving a "
       "start-of-session and end-of-session comparison.")

h1("7. Task performance data")
para("One of two 'base' tasks is assigned per session:")
bullet("Fitts' Law (pointing)", "target size and distance, movement time, click accuracy/"
       "error, and throughput per trial. Numeric performance only.")
bullet("Typing", "the fixed prompt sentence shown and what the participant typed in response, "
       "plus derived typing speed (WPM), error rate, inter-key timing, and backspace count. "
       "The prompts are fixed, neutral, non-personal sentences from a pre-defined corpus - "
       "not the participant's own free-form writing.")
para("One of two 'fatigue' tasks is assigned per session:")
bullet("Cognitive (Stroop / AX-CPT)", "the stimulus shown, the participant's response, "
       "whether it was correct, reaction time, and whether they timed out. Fixed words/"
       "colours/letters defined by the task, not personal content.")
bullet("Physical", "active duration, paused duration, number of pauses, and whether the "
       "protocol was completed - timing/completion data about a researcher-approved light "
       "movement protocol, not biometric data itself (see Section 8).")

h1("8. Physiological / device measurements")
para("When applicable, a researcher enters ECG readings (heart rate, HRV, and any additional "
     "values a device provides) for a participant, tagged to the exact session and task block "
     "it was recorded during - either from the project's Raspberry Pi sensor or, going "
     "forward, from a lab-grade ECG device used in parallel with the software test.")

h1("9. Session event log")
para("A timestamped internal log of protocol milestones (consent given, calibration "
     "completed, each rating recorded, a break skipped, the session completed or withdrawn) - "
     "used to confirm data integrity and reconstruct the session timeline, not to capture "
     "additional personal content.")

h1("10. Withdrawal")
para("A participant may withdraw at any time; this stops all further collection for that "
     "session and the session is marked accordingly. Already-collected data is handled under "
     "the retention policy stated at consent.")

h1("11. Planned, not yet active: phone-use study")
para("A separate, optional study (on hold, not currently collecting data) is designed to "
     "record - via a dedicated Android app, once built - which apps a participant uses, for "
     "how long, how often they are opened, and how much scrolling occurs, without recording "
     "any content viewed. Listed here for completeness.")

h1("12. What is explicitly NOT collected")
bullet("No raw or recorded video/audio", "only the derived numeric attention signal (Section "
       "5), plus optional low-resolution still photos if separately consented.")
bullet("No content typed outside the fixed task prompts", "no free-form personal writing, "
       "messages, or files.")
bullet("No browsing history, no data from outside the study session.")
bullet("No participant name/email is ever attached to a research record", "only the "
       "separate, admin-only contacts table (Section 1) holds that link.")

h1("13. Where this lives and who can see it")
para("All data is stored in a single database with row-level security: the participant's "
     "device can only ever insert its own data, never read anyone else's; only an approved, "
     "admin-authorised researcher account can read research data; contact details and camera "
     "images are further restricted to admin accounts only. Every data export and every "
     "access-approval decision is recorded in an append-only audit log.")

doc.save(OUT)
print("wrote", OUT)
