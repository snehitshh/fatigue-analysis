"""Generate FatigueIDPro_Model_Comparison_Study.docx from dl-model/results.
Run (after the dl-model pipeline): python scripts/make-model-comparison-doc.py
"""
from datetime import date
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x0F, 0x2B, 0x4A); BLUE = RGBColor(0x1D, 0x4E, 0xD8); GREY = RGBColor(0x47, 0x55, 0x69)
ROOT = "e:/Data-center/fatigue-analysis"
DL = f"{ROOT}/dl-model"

doc = Document()
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10.5)


def title(t, sub):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(sub); r2.font.size = Pt(12); r2.font.color.rgb = BLUE
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Phase P0 + P1  |  " + date.today().strftime("%d %B %Y"))
    r3.font.size = Pt(9); r3.font.color.rgb = GREY


def h1(t):
    doc.add_paragraph(); h = doc.add_heading(level=1)
    r = h.add_run(t); r.font.color.rgb = NAVY; r.font.size = Pt(15)


def h2(t):
    h = doc.add_heading(level=2); r = h.add_run(t); r.font.color.rgb = BLUE; r.font.size = Pt(12)


def para(t, grey=False):
    p = doc.add_paragraph(); r = p.add_run(t)
    if grey: r.font.color.rgb = GREY; r.font.size = Pt(10)


def bullet(t):
    doc.add_paragraph(t, style="List Bullet")


def add_results_table(csv_path):
    results = pd.read_csv(csv_path)
    rows = [["Model", "MAE", "RMSE", "R2"]] + [
        [r["model"], f"{r['MAE_mean']:.3f} +/- {r['MAE_std']:.3f}",
         f"{r['RMSE_mean']:.3f} +/- {r['RMSE_std']:.3f}", f"{r['R2_mean']:.3f}"]
        for _, r in results.iterrows()
    ]
    t = doc.add_table(rows=len(rows), cols=4); t.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.text = str(val)
            for pp in c.paragraphs:
                for rn in pp.runs:
                    rn.font.size = Pt(9)
                    if ri == 0: rn.bold = True
    return results


title("Fatigue Model Comparison Study",
      "Classic ML (P0) vs a lightweight deep model (P1), across two modalities")
para("Goal: establish the classic-ML benchmark (Phase P0) and a first lightweight deep "
     "model (Phase P1) per signal modality, using the datasets parsed so far - FatigueSet "
     "(ECG/HRV) and a newly added Mendeley EMG dataset. Not the final model - a reference "
     "point the later, larger model (trained once our own collection has real data) must beat.",
     grey=True)

# ---------------------------------------------------------------- ECG
h1("1. ECG / HRV — FatigueSet (12 subjects, 677 windows)")
para("Features: hr, rmssd, sdnn, lf_hf (median-imputed where a 30s window makes LF/HF "
     "unreliable), sss_pretask, gvas_sleepy, intensity_level. Target: physical fatigue "
     "rating (continuous, ~0-100). Validation: Leave-One-Subject-Out (12 folds).")
h2("1a. Classic ML benchmark (P0)")
ecg_p0 = add_results_table(f"{DL}/results/model_comparison_ecg_fatigueset.csv")
h2("1b. Lightweight deep model (P1) vs P0")
dl_ecg = pd.read_csv(f"{DL}/results/dl_model_ecg_fatigueset.csv").iloc[0]
best_ecg = ecg_p0.iloc[0]
para(f"MLP (PyTorch): MAE {dl_ecg['MAE_mean']:.3f}, RMSE {dl_ecg['RMSE_mean']:.3f}, "
     f"R2 {dl_ecg['R2_mean']:.3f}. Best classic ({best_ecg['model']}): MAE {best_ecg['MAE_mean']:.3f}.")
para("Classic ML wins on this modality. Even the Dummy (mean) baseline scores negative R2 "
     "under LOSO: fatigue baselines vary a lot between subjects, so a population model needs "
     "per-subject calibration more than a fancier algorithm at 12 subjects.")

# ---------------------------------------------------------------- EMG
h1("2. EMG — Mendeley 'Muscle Fatigue in Biceps and Triceps' (30 subjects, 480 windows)")
para("Newly added dataset (DOI 10.17632/8j2p29hnbv.1), freely downloadable. Features: "
     "RMS, MAV, waveform length, zero-crossings, slope-sign-changes (classical time-domain "
     "EMG fatigue indicators - frequency-domain MDF/MNF were skipped because the source "
     "metadata does not reliably state the sampling rate). Label: a rep-ordinal fatigue "
     "proxy (each recording has 4 marked repetitions; rep 1 = freshest, rep 4 = most "
     "fatigued - a standard assumption for this protocol, documented as a proxy not a "
     "physiological ground truth). Validation: Leave-One-Subject-Out (30 folds).")
h2("2a. Classic ML benchmark (P0)")
emg_p0 = add_results_table(f"{DL}/results/model_comparison_emg_mendeley.csv")
h2("2b. Lightweight deep model (P1) vs P0")
dl_emg = pd.read_csv(f"{DL}/results/dl_model_emg_mendeley.csv").iloc[0]
best_emg = emg_p0.iloc[0]
para(f"MLP (PyTorch): MAE {dl_emg['MAE_mean']:.3f}, RMSE {dl_emg['RMSE_mean']:.3f}, "
     f"R2 {dl_emg['R2_mean']:.3f}. Best classic ({best_emg['model']}): MAE {best_emg['MAE_mean']:.3f}, "
     f"R2 {best_emg['R2_mean']:.3f}.")
para("Classic ML (Random Forest) wins here too, but by a smaller margin. Unlike ECG, this "
     "modality shows a genuinely positive R2 - RMS/MAV/waveform-length track muscle fatigue "
     "well even across 30 different people, because these are direct physical exertion "
     "signals rather than a subjective/physiological state that varies by personality and "
     "baseline arousal.")

# ---------------------------------------------------------------- interpretation
h1("3. Interpretation")
bullet("Classic ML beats a from-scratch deep model on BOTH modalities at this sample size "
       "(a few hundred rows). This matches well-established findings on small tabular "
       "datasets - deep networks need substantially more data before they out-perform tree "
       "ensembles or kernel methods; it is an expected result, not a failure.")
bullet("EMG generalises across subjects much better than ECG (positive vs negative R2). "
       "Physical/muscular fatigue signals are more directly measurable than the "
       "person-dependent physiological/subjective state ECG-derived fatigue partly reflects.")
bullet("Recommendation: normalise labels/features per subject before pooling; prioritise "
       "collecting more subjects (our own data collection, P2) over enlarging the network; "
       "revisit the deep-model comparison once our multimodal data is available.")

# ---------------------------------------------------------------- verification model
h1("4. Data-quality verification model (the primary purpose of this work)")
para("An unsupervised IsolationForest, one per modality, fit on the valid feature "
     "distribution (no fatigue label needed). At real collection time it scores each new "
     "incoming window; an implausible reading (bad electrode contact, motion artefact, "
     "sensor fault) is flagged for the operator instead of being silently kept.")
bullet("Sanity-checked: a physiologically plausible ECG reading (HR 75, normal HRV) scores "
       "as normal; an implausible one (HR 2, HRV 5000ms) is correctly flagged as an anomaly.")
bullet("~5% of the existing data is flagged in each modality (by design - the contamination "
       "rate is a tunable parameter).")
bullet("Can be extended with new data with no manual labelling required (unsupervised) - see "
       "Section 5.")

# ---------------------------------------------------------------- finetune
h1("5. Continual-learning / fine-tune mechanism")
para("A scaffold that continues training a saved model on new same-schema data (low "
     "learning rate, reusing the ORIGINAL fitted preprocessor so feature scaling stays "
     "consistent) and refits the quality detector on old+new combined. Proven with a "
     "running self-check against a real held-out slice of the existing data. Not yet run "
     "on real field data - our own collection currently has ~0 sessions in the live "
     "database; this activates once it does.")

# ---------------------------------------------------------------- next
h1("6. Next steps")
bullet("Parse additional datasets/modalities as they become available (video/PERCLOS next).")
bullet("P2: fine-tune/calibrate the deep model on our own multimodal sessions (KSS + Borg + "
       "NASA-TLX + ECG + task performance) once real collection data exists.")
bullet("Re-run this exact comparison at that point - the expectation is the deep model "
       "should close the gap with more, richer, multimodal data.")

doc.save(f"{ROOT}/FatigueIDPro_Model_Comparison_Study.docx")
print("wrote FatigueIDPro_Model_Comparison_Study.docx")
