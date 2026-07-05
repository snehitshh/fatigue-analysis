"""Generate FatigueIDPro_Phone_App_Architecture.docx - explains the Android app for
the phone-usage study: what it collects, permissions, why it's difficult, and whether
sideloading (Chrome APK download) can collect data. Run: python scripts/make-app-doc.py
"""
from datetime import date
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn

NAVY = RGBColor(0x0F, 0x2B, 0x4A)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GREY = RGBColor(0x47, 0x55, 0x69)
GREEN = RGBColor(0x15, 0x80, 0x3D)
RED = RGBColor(0xB4, 0x1E, 0x1E)

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)


def add_hyperlink(p, url, text):
    r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), r_id)
    run = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '1D4ED8'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    run.append(rPr); t = OxmlElement('w:t'); t.text = text; run.append(t); h.append(run); p._p.append(h)


def title(text, sub):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(sub); r2.font.size = Pt(12); r2.font.color.rgb = BLUE
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Technical note for review  |  " + date.today().strftime("%d %B %Y"))
    r3.font.size = Pt(9); r3.font.color.rgb = GREY


def h1(text):
    doc.add_paragraph()
    h = doc.add_heading(level=1); r = h.add_run(text); r.font.color.rgb = NAVY; r.font.size = Pt(15)


def h2(text):
    h = doc.add_heading(level=2); r = h.add_run(text); r.font.color.rgb = BLUE; r.font.size = Pt(12)


def para(text, color=None, bold=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold
    if color is not None: r.font.color.rgb = color
    return p


def bullet(name, desc):
    p = doc.add_paragraph(style="List Bullet")
    if name:
        r = p.add_run(name + " — "); r.bold = True; r.font.color.rgb = NAVY
    p.add_run(desc)


# ---------------------------------------------------------------- cover
title("FatigueIDPro — Phone-Usage Study App",
      "Android architecture, permissions, distribution, and feasibility")
para("This note explains the mobile app needed for the phone-usage (scrolling) part of the study: what it "
     "records, the Android permissions it needs, why it is technically difficult, and — the key question "
     "— whether distributing the app by direct download (sideloading via a Chrome link) instead of the "
     "Google Play Store still lets us collect data.", GREY)

# ---------------------------------------------------------------- 1
h1("1. What the study needs to measure")
para("We want to measure real phone use as a source of fatigue: how much a participant uses and scrolls "
     "real social apps (Instagram, Facebook, YouTube, etc.) on their own phone, and how that activity drifts "
     "over a chosen window (30 / 60 / 120 minutes). Concretely, per app:")
bullet("Foreground time", "how long each app is actively on screen.")
bullet("Opens", "how many times each app is launched/returned to.")
bullet("Scroll-event count", "how many scroll gestures occur in each app (a proxy for scrolling intensity).")
para("We do NOT capture any content — not what they read, type, or view. Only which app, how long, how "
     "often, and how much scrolling. A KSS sleepiness rating is taken before and after.", GREY)

# ---------------------------------------------------------------- 2
h1("2. Why a normal website / browser cannot do this")
para("Phones sandbox every app: from a website or an ordinary app you cannot see what happens inside "
     "Instagram or Facebook. So a browser-based page can never read another app's usage or scrolling. This "
     "must be a native Android app that the participant installs and explicitly grants special permissions to. "
     "(Our web page already handles registration and the KSS ratings; the native app does the recording.)")

# ---------------------------------------------------------------- 3
h1("3. How the app works (architecture)")
bullet("Capacitor shell", "the existing FatigueIDPro web app is wrapped as an Android app (Capacitor), so the "
       "registration + KSS screens are reused unchanged.")
bullet("Native plugin", "a small Android module exposes window.FatiguePhoneUsage.start()/stop() to the web "
       "layer. It runs a foreground service that records usage and uploads it.")
bullet("Two Android data sources", "UsageStatsManager (time on app + opens) and an AccessibilityService "
       "(foreground app + scroll-event counts).")
bullet("Same database", "the app uploads per-app interval rows and session totals to the same Supabase "
       "backend already built for this (phone_usage_sessions + app_usage_intervals).")

# ---------------------------------------------------------------- 4
h1("4. Permissions the participant must grant")
para("These are all granted by the participant, once, in Android Settings after installing. None of them can "
     "be enabled silently.", GREY)

rows = [
    ("Permission", "What it enables", "How it's granted"),
    ("Usage Access\n(PACKAGE_USAGE_STATS)", "Per-app foreground time and launch counts (UsageStatsManager).",
     "Settings > Apps > Special access > Usage access > enable for FatigueIDPro."),
    ("Accessibility Service", "Detects the foreground app and scroll events (TYPE_VIEW_SCROLLED) across other apps.",
     "Settings > Accessibility > FatigueIDPro > turn on. (This is the 'record over other apps' capability.)"),
    ("Foreground service +\nnotification", "Keeps recording reliably while the phone is used/locked; shows a persistent 'recording' notice.",
     "Granted at install / first run; a notification stays visible while active."),
    ("Ignore battery optimisation", "Stops the phone from killing the recorder in the background.",
     "A one-tap prompt the app shows; recommended on Xiaomi/Oppo/Samsung etc."),
]
table = doc.add_table(rows=len(rows), cols=3)
table.style = "Light Grid Accent 1"
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = table.cell(ri, ci); cell.text = val
        for pph in cell.paragraphs:
            for rn in pph.runs:
                rn.font.size = Pt(8.5)
                if ri == 0: rn.bold = True

# ---------------------------------------------------------------- 5
h1("5. The key question: Play Store vs direct download (sideload)")
para("Planned approach: distribute the app as an APK by direct download (e.g. a Chrome link), not through "
     "the Google Play Store.", bold=True)

h2("Will sideloading still collect data? Yes.")
para("The difficulty with this app is a GOOGLE PLAY POLICY, not an Android limitation. Google Play forbids "
     "most apps from using the Accessibility Service for anything other than helping users with disabilities; "
     "a research usage-tracker would very likely be REJECTED from the Play Store. That policy only applies to "
     "apps distributed THROUGH the Play Store.")
para("If we hand out the APK directly (sideload), we are not subject to Play review, so that restriction does "
     "not block us. The Android operating system itself fully supports Usage Access and the Accessibility "
     "Service for sideloaded apps — the permissions work exactly the same. So: sideloading is not a "
     "workaround that loses data; it is the correct distribution method for a consented research app that "
     "needs these permissions.", GREEN)

h2("The catch with sideloading (all manageable):")
bullet("'Unknown sources' warning", "the participant must allow installing the APK from the browser/files app "
       "(a one-time Settings toggle). Play Protect may show a warning they tap through.")
bullet("Android 13+ 'restricted settings'", "for sideloaded apps, the Accessibility toggle is greyed out "
       "until the participant opens App info > (3-dot menu) > 'Allow restricted settings'. One extra step, but "
       "expected and documented.")
bullet("No auto-updates", "we distribute new versions by sending a new link; there is no Play auto-update.")
bullet("Manual permission setup", "each participant must enable Usage Access + Accessibility themselves; we "
       "provide an in-app step-by-step screen that opens the right Settings page.")

# ---------------------------------------------------------------- 6
h1("6. Why it's difficult (summary of risks)")
bullet("Play Store would reject it", "hence sideloading — which is fine for a consented research study, "
       "but means it is not a public app-store product.")
bullet("Permission friction", "3-4 manual permission steps per participant; some will need help. An in-app "
       "guided setup mitigates this.")
bullet("OEM battery killers", "Xiaomi/Oppo/Vivo/Samsung aggressively kill background services; we must ask "
       "participants to disable battery optimisation, and even then long windows can be interrupted.")
bullet("Android only", "iOS has no equivalent public API to read other apps' usage or scrolling (Apple's "
       "Screen Time is locked behind Family Controls and does not expose scroll events). This study is "
       "effectively Android-only.")
bullet("Trust & consent", "the Accessibility permission is powerful and participants are rightly cautious; "
       "clear, honest consent (we record only app + time + scroll counts, never content) is essential.")

h2("Fallback if Accessibility proves too hard")
para("If the Accessibility permission is a blocker for some participants or devices, Usage Access ALONE "
     "(time-on-app + opens, no scroll counts) still gives a strong phone-fatigue signal and needs only one, "
     "less-sensitive permission. We can run in that reduced mode and still analyse fatigue from time and "
     "open-frequency drift.")

# ---------------------------------------------------------------- 7
h1("7. What is already built vs what remains")
bullet("Done", "database (phone_usage_sessions + app_usage_intervals, security, export views, finalize "
       "function — live); web registration + KSS flow; the native bridge contract "
       "(window.FatiguePhoneUsage) the app must implement; admin console shows the data.")
bullet("Remaining", "the native Android plugin (UsageStats reader + AccessibilityService + foreground service "
       "+ background upload), the in-app permission-setup screens, and packaging the APK for download.")

# ---------------------------------------------------------------- refs
h1("References")
p = doc.add_paragraph(style="List Bullet"); add_hyperlink(p, "https://developer.android.com/reference/android/app/usage/UsageStatsManager", "Android UsageStatsManager (per-app usage)")
p = doc.add_paragraph(style="List Bullet"); add_hyperlink(p, "https://developer.android.com/guide/topics/ui/accessibility/service", "Android AccessibilityService (scroll / foreground events)")
p = doc.add_paragraph(style="List Bullet"); add_hyperlink(p, "https://support.google.com/googleplay/android-developer/answer/10964491", "Google Play policy: Accessibility API usage")

out = "FatigueIDPro_Phone_App_Architecture.docx"
doc.save(out)
print("wrote", out)
